"""Turn uploaded files into plain text. Everything here is local and free."""

import csv
import io
import json
import shutil
from dataclasses import dataclass
from pathlib import PurePath

TEXT_EXT = {".txt", ".md", ".markdown"}
JSON_EXT = {".json"}
PDF_EXT = {".pdf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm", ".avi"}
TABLE_EXT = {".csv", ".xlsx"}

MAX_TABLE_ROWS = 5000


class IngestError(ValueError):
    """A problem with the uploaded content that the user can act on."""


class ToolMissing(IngestError):
    """A system tool (Tesseract, ffmpeg) needed for this file type is not installed."""


def detect_kind(filename: str) -> str:
    """text | json | pdf | image | audio | video | table"""
    ext = PurePath(filename).suffix.lower()
    for kind, exts in (
        ("text", TEXT_EXT), ("json", JSON_EXT), ("pdf", PDF_EXT), ("image", IMAGE_EXT),
        ("audio", AUDIO_EXT), ("video", VIDEO_EXT), ("table", TABLE_EXT),
    ):
        if ext in exts:
            return kind
    raise IngestError(f"Unsupported file type '{ext or filename}'")


def default_source_type(kind: str) -> str:
    return {"text": "transcript", "json": "transcript", "pdf": "pdf", "image": "ocr",
            "audio": "audio", "video": "video", "table": "survey"}[kind]


def decode_text(data: bytes) -> str:
    # UTF-16 only with a byte-order mark: without one, arbitrary Latin-1 bytes can
    # happen to decode as valid UTF-16 and come out as garbage.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeError:
        return data.decode("latin-1")


def _json_strings(node, out: list[str]) -> None:
    if isinstance(node, str):
        if node.strip():
            out.append(node)
    elif isinstance(node, dict):
        # a chat/transcript turn like {"speaker": "...", "text": "..."}
        speaker = next((node[k] for k in ("speaker", "role", "author", "name") if isinstance(node.get(k), str)), None)
        body = next((node[k] for k in ("text", "content", "message", "body") if isinstance(node.get(k), str)), None)
        if body is not None:
            out.append(f"{speaker}: {body}" if speaker else body)
        else:
            for value in node.values():
                _json_strings(value, out)
    elif isinstance(node, list):
        for item in node:
            _json_strings(item, out)


def parse_json(data: bytes) -> str:
    try:
        parsed = json.loads(decode_text(data))
    except json.JSONDecodeError as exc:
        raise IngestError(f"Invalid JSON: {exc.msg}") from exc
    lines: list[str] = []
    _json_strings(parsed, lines)
    return "\n".join(lines)


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except (PyPdfError, ValueError) as exc:
        raise IngestError(f"Could not read the PDF: {exc}") from exc
    text = "\n\n".join(p for p in pages if p.strip())
    if not text.strip():
        raise IngestError(
            "This PDF has no selectable text (it looks scanned). Upload its pages as images to use OCR."
        )
    return text


def parse_image(data: bytes) -> str:
    import pytesseract
    from PIL import Image, UnidentifiedImageError

    if not shutil.which("tesseract"):
        raise ToolMissing("Tesseract OCR is not installed on the server, so images cannot be read.")
    try:
        text = pytesseract.image_to_string(Image.open(io.BytesIO(data)))
    except UnidentifiedImageError as exc:
        raise IngestError("Could not read the image file") from exc
    if not text.strip():
        raise IngestError("No text was found in the image")
    return text


@dataclass
class TableGroup:
    segment: str | None  # value of the segment column, if one was chosen
    text: str  # "Row N: ..." lines


def _read_rows(data: bytes, filename: str) -> list[list[str]]:
    if PurePath(filename).suffix.lower() == ".csv":
        text = decode_text(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    from openpyxl import load_workbook

    try:
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        return [["" if c is None else str(c).strip() for c in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise IngestError(f"Could not read the spreadsheet: {exc}") from exc


def _pick_column(headers: list[str], wanted: str | None, rows: list[list[str]]) -> int:
    if wanted:
        lowered = [h.lower() for h in headers]
        if wanted.lower() not in lowered:
            raise IngestError(f"Column '{wanted}' not found. Columns: {', '.join(headers)}")
        return lowered.index(wanted.lower())
    sample = rows[:200]  # default: the column with the longest average text
    averages = [
        sum(len(r[i]) for r in sample if i < len(r)) / max(len(sample), 1) for i in range(len(headers))
    ]
    return max(range(len(headers)), key=lambda i: averages[i])


def parse_table(
    data: bytes, filename: str, text_column: str | None = None, segment_column: str | None = None
) -> list[TableGroup]:
    """Survey or ticket export -> one group per segment value (or one group overall)."""
    rows = [r for r in _read_rows(data, filename) if any(r)]
    if len(rows) < 2:
        raise IngestError("The file needs a header row and at least one data row")
    headers, body = rows[0], rows[1:MAX_TABLE_ROWS + 1]
    text_idx = _pick_column(headers, text_column, body)
    seg_idx = _pick_column(headers, segment_column, body) if segment_column else None

    groups: dict[str | None, list[str]] = {}
    for number, row in enumerate(body, start=2):
        text = row[text_idx] if text_idx < len(row) else ""
        if not text:
            continue
        segment = (row[seg_idx] if seg_idx is not None and seg_idx < len(row) else "") or None
        groups.setdefault(segment, []).append(f"Row {number}: {text}")
    if not groups:
        raise IngestError(f"Column '{headers[text_idx]}' has no text")
    return [TableGroup(segment, "\n".join(lines)) for segment, lines in groups.items()]
